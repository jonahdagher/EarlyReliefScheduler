
from openpyxl import load_workbook
import re

from sqlalchemy import *
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import declarative_base, sessionmaker, relationship



class rawScheduleEntry():
    def __init__(self, date:str, value:str, color:str = None):
        self.date = date
        self.value = value
        self.color = color

    def __repr__(self):
        return f"date: {self.date} | value: {self.value} | color: {self.color}"

class providerSchedule():
    def __init__(self, name:str):
        self.name = name
        self.schedule = []

    def __iter__(self):
        return iter(self.schedule)

    def add_entry(self, entry: rawScheduleEntry):
        if not isinstance(entry, rawScheduleEntry):
            raise ValueError("entries should be of type <rawScheduleEntry>")
        else:
            self.schedule.append(entry)


class Schedule():
    def __init__(self, path, name_col = 1, date_row = 3, revsion_index = (1,2)):
        self.path = path
        self.data = self.read()

        self.name_col = name_col
        self.date_row = date_row
        self.revision_index = revsion_index

        self.bottom_bound = self.getBottomBound()
        self.right_bound = self.getRightBound()

        self.year_month = self.getRevisedDate()[:7]

        self.all_colors = {
            "rgb": set(),
            "indexed": set(),
            "theme": set()
        }

    def getRevisedDate(self):
        revision_cell_value = self.data.cell(column=self.revision_index[0], row=self.revision_index[1]).value
        cleaned_revision_cell_value = re.sub(r"[^0-9/]", "", revision_cell_value)
        day_month_year = cleaned_revision_cell_value.split('/')

        return f"{day_month_year[2]}-{day_month_year[0]}-{day_month_year[1]}"
    
    def getYearMonth(self):
        return self.getRevisedDate()[:7]

    def read(self):
        wb = load_workbook(self.path, data_only=True)
        self.data = wb.active
        return wb.active
    
    def getBottomBound(self):
        
        current_max = self.date_row+1

        #Yields cells from name colum starting under the date row
        name_column = self.data.iter_cols(
            min_col=self.name_col, 
            max_col=self.name_col, 
            min_row=self.date_row+1) 
        
        for cell_tuple in name_column:
            for cell in cell_tuple:
                if "#" in cell.value:
                    self.bottom_bound = current_max
                    return current_max - 1
                else:
                    current_max += 1

        return self.data.max_row
    
    def getRightBound(self):
        current_max = self.name_col

        date_row = self.data.iter_rows(
            min_col=self.name_col+1, 
            min_row=self.date_row,
            max_row=self.date_row)
        
        for cell_tuple in date_row:
            for cell in cell_tuple:
                if str(cell.value).isdigit():
                    current_max += 1
                else:
                    return current_max
        return self.data.max_column

    
    def getProviderRows(self):

        providers = []

        name_column = self.data.iter_cols(
            min_col=self.name_col, 
            max_col=self.name_col, 
            min_row=self.date_row+1,
            max_row=self.bottom_bound)
        
        for cell_tuple in name_column:
            for cell in cell_tuple:
                providers.append({"name":cell.value, "row":cell.row})

        return providers
    
    
    def getProviderSchedules(self):

        schedules = dict()

        for provider in self.getProviderRows():

            provider_schedule =  providerSchedule(provider["name"])

            provider_row = self.data.iter_rows(
                min_col=self.name_col+1,
                max_col=self.right_bound,
                min_row=provider["row"],
                max_row=provider["row"])

            for cell_tuple in provider_row:
                for cell in cell_tuple:
                    cell_date = self.year_month+"-"+f"{self.getAssociatedDate(cell.column)['date']:02}"
                    cell_data = rawScheduleEntry(cell_date, cell.value, self.getCellColor(cell))
                    provider_schedule.add_entry(cell_data)
            
            schedules[provider["name"]] = provider_schedule
        return schedules


    def getAssociatedDate(self, col:int) -> dict:
        '''
        Docstring for getAssociatedDate
        :param col: col of the cell to get associated date from
        :type col: int
        :return: returns the date and day associated with a cell
        :rtype: dict
        '''

        ALLOWED_DAYS = set(["MON","TUE","WED","THU","FRI"])

        date = self.data.cell(row=self.date_row, column=col).value
        day = self.data.cell(row=self.date_row-1, column=col).value
        if day not in ALLOWED_DAYS:
            raise ValueError(f"{col, self.date_row}Attempting to access day {day} is not in {ALLOWED_DAYS}")
        return {"date":date, "day":day}
    
    def getCellColor(self, cell):
        fill = cell.fill

        if not fill or fill.fill_type != "solid":
            return None
        
        color = fill.start_color

        if color is None:
            return None
        
        if color.type == "rgb" and color.rgb:
            self.all_colors[color.type].add(color.rgb[-6:])
            return "#" + color.rgb[-6:] #remove alpha values
        
        else:
            return "#808080"
        
import pandas as pd

import pandas as pd

def times_days_to_df(
    time_dict: dict,
    times=("330", "530", "730", "930"),
    days_in_order=None,
    days_per_row=5,
    day_spacer_col=True,
    block_spacer_row=True,
):

    # infer days if not provided: union of all time_dict[time] keys
    if days_in_order is None:
        day_set = set()
        for t in times:
            day_set |= set(time_dict.get(t, {}).keys())
        days_in_order = sorted(day_set)

    day_panels = []
    for day in days_in_order:
        # Build one day panel: 4 columns (one per time)
        cols = {}
        max_len = 0

        # pull lists + compute max height for this day
        for t in times:
            items = time_dict.get(t, {}).get(day, [])
            if items is None:
                items = []
            if not isinstance(items, list):
                items = [str(items)]
            cols[t] = items
            max_len = max(max_len, len(items))

        # pad all time columns to same length
        for t in times:
            cols[t] = cols[t] + [""] * (max_len - len(cols[t]))

        panel = pd.DataFrame(cols)

        # rename columns to include day label so they’re unique
        panel.columns = [f"{day} {t}" for t in times]

        # optional spacer column after each day panel
        if day_spacer_col:
            panel[f"{day} _"] = [""] * len(panel)

        day_panels.append(panel)

    # now pack 5 day-panels per horizontal block, then stack blocks vertically
    blocks = []
    for i in range(0, len(day_panels), days_per_row):
        block = pd.concat(day_panels[i:i+days_per_row], axis=1)
        blocks.append(block)

    if not blocks:
        return pd.DataFrame()

    out = blocks[0]
    for b in blocks[1:]:
        if block_spacer_row:
            blank_row = pd.DataFrame({c: [""] for c in out.columns})
            out = pd.concat([out, blank_row, b], ignore_index=True)
        else:
            out = pd.concat([out, b], ignore_index=True)

    return out

